"""Semantic and hybrid company-document retrieval backed by PostgreSQL/pgvector.

Structured business facts remain in the existing SQLite demo database and are
queried with SQL. This module embeds document/review passages through Featherless,
stores vectors and provenance in PostgreSQL, and combines vector similarity with
PostgreSQL full-text search when an analyst explicitly requests retrieval.
"""
from __future__ import annotations

import logging
from collections import defaultdict
from contextlib import contextmanager
from contextvars import ContextVar
from hashlib import sha256
from io import BytesIO
import json
import math
import os
from pathlib import Path
import re
import sqlite3
from typing import Any, Iterable
from urllib.parse import parse_qsl, unquote, urlsplit
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from db_diagnostics import log_database_error
from public_safety import reserve_public_request

from jev import load_local_environment
from observability import trace_component

ROOT = Path(__file__).resolve().parent
BUSINESS_DB_PATH = ROOT / "demo_company" / "data" / "business_demo.sqlite3"
DEFAULT_DOCS_PATH = ROOT / "demo_company" / "docs"
DEFAULT_UPLOADS_PATH = ROOT / "demo_company" / "uploads"
MAX_CHUNK_CHARS = 1800
CHUNK_OVERLAP = 180
EMBEDDING_DIMENSIONS = 1024
EMBEDDING_BATCH_SIZE = 32
DEFAULT_EMBEDDING_MODEL = "Qwen/Qwen3-Embedding-8B"
EXCLUDED_DOCS = {"scenario_truth.json"}
SUPPORTED_EXTENSIONS = {".md", ".txt", ".json", ".sql", ".pdf", ".docx"}
TENANT_ENV = "RAG_TENANT_ID"
_LOGGER = logging.getLogger(__name__)
_ACTIVE_RETRIEVAL_TENANTS: ContextVar[tuple[str, ...] | None] = ContextVar("active_retrieval_tenants", default=None)


@contextmanager
def retrieval_tenant_scope(tenant_ids: list[str] | tuple[str, ...] | None):
    """Scope implicit specialist retrieval to the current Streamlit session."""
    normalized = tuple(value for value in (tenant_ids or ()) if isinstance(value, str) and value.strip())
    token = _ACTIVE_RETRIEVAL_TENANTS.set(normalized or None)
    try:
        yield
    finally:
        _ACTIVE_RETRIEVAL_TENANTS.reset(token)


class RagError(RuntimeError):
    """Raised when retrieval, indexing, or document parsing cannot complete."""


def _tenant_id(tenant_id: str | None = None) -> str:
    if tenant_id is None:
        load_local_environment()
        value = os.getenv(TENANT_ENV, "demo").strip()
    else:
        value = tenant_id.strip()
    if not value or len(value) > 128:
        raise RagError("RAG_TENANT_ID must be a non-empty value no longer than 128 characters.")
    return value


def _database_url() -> str:
    load_local_environment()
    value = os.getenv("RAG_DATABASE_URL", "").strip()
    if not value:
        raise RagError(
            "Semantic retrieval needs PostgreSQL with pgvector. Set RAG_DATABASE_URL "
            "in .env and start the local vector database; see the project README."
        )
    return value


def _embedding_model() -> str:
    load_local_environment()
    return os.getenv("FEATHERLESS_EMBEDDING_MODEL", DEFAULT_EMBEDDING_MODEL).strip() or DEFAULT_EMBEDDING_MODEL


def _connect():
    try:
        import psycopg
    except ImportError as exc:
        raise RagError("Install the project dependencies with: py -m pip install -r requirements.txt") from exc
    database_url = _database_url()
    parsed_url = urlsplit(database_url)
    host = (parsed_url.hostname or "").casefold()
    sslmode = "prefer" if host in {"localhost", "127.0.0.1", "::1"} else "require"
    try:
        # Decode URI components before passing credentials separately. This
        # keeps passwords with reserved URL characters equivalent to the
        # explicit host/user/password connection used by deployment checks.
        connection_options = dict(parse_qsl(parsed_url.query, keep_blank_values=True))
        connection_options.update(
            {
                "host": parsed_url.hostname,
                "port": parsed_url.port or 5432,
                "dbname": unquote(parsed_url.path.lstrip("/")),
                "user": unquote(parsed_url.username or ""),
                "password": unquote(parsed_url.password or ""),
                "sslmode": sslmode,
            }
        )
        # Supabase transaction poolers do not support prepared statements.
        return psycopg.connect(
            **connection_options,
            connect_timeout=5,
            prepare_threshold=None,
        )
    except Exception as exc:
        # Log only safe classification fields; raw driver text may expose connection details.
        parsed_url = urlsplit(database_url)
        username = unquote(parsed_url.username or "")
        username_hint = (
            f"{username.split('.', 1)[0]}.…{username[-4:]}"
            if "." in username and len(username) > 4
            else (f"…{username[-4:]}" if len(username) > 4 else "other")
        )
        _LOGGER.error(
            "RAG database target: source=%s host=%s port=%s database=%s username_hint=%s",
            os.getenv("_RAG_DATABASE_URL_SOURCE", "environment_or_dotenv"),
            parsed_url.hostname or "unknown",
            parsed_url.port or "default",
            parsed_url.path.lstrip("/") or "unknown",
            username_hint,
        )
        log_database_error(_LOGGER, "rag_connect", exc)
        raise RagError(
            "Could not connect to the RAG PostgreSQL database. Check RAG_DATABASE_URL "
            "and start the pgvector service."
        ) from exc


def _ensure_schema(connection: Any) -> None:
    try:
        from psycopg import sql

        connection.execute("CREATE EXTENSION IF NOT EXISTS vector")
        extension_schema = connection.execute(
            """SELECT namespace.nspname
               FROM pg_extension AS extension
               JOIN pg_namespace AS namespace ON namespace.oid = extension.extnamespace
               WHERE extension.extname = 'vector'"""
        ).fetchone()
        if not extension_schema:
            raise RuntimeError("pgvector extension is not installed")
        # Supabase commonly installs pgvector under `extensions`; local Docker
        # installs it under `public`. Keep the path transaction-local for poolers.
        connection.execute(
            sql.SQL("SET LOCAL search_path TO public, {}").format(sql.Identifier(str(extension_schema[0])))
        )
        connection.execute(
            f"""CREATE TABLE IF NOT EXISTS rag_chunks (
                id BIGSERIAL PRIMARY KEY,
                tenant_id TEXT NOT NULL,
                source_type TEXT NOT NULL,
                source_name TEXT NOT NULL,
                source_ref TEXT NOT NULL,
                title TEXT NOT NULL,
                section TEXT NOT NULL,
                content TEXT NOT NULL,
                content_hash TEXT NOT NULL,
                embedding vector({EMBEDDING_DIMENSIONS}) NOT NULL,
                search_tsv tsvector GENERATED ALWAYS AS (
                    setweight(to_tsvector('english'::regconfig, coalesce(title, '')), 'A') ||
                    setweight(to_tsvector('english'::regconfig, coalesce(section, '')), 'B') ||
                    setweight(to_tsvector('english'::regconfig, coalesce(content, '')), 'C')
                ) STORED,
                updated_at TIMESTAMPTZ NOT NULL DEFAULT now(),
                UNIQUE (tenant_id, source_name, source_ref)
            )"""
        )
        connection.execute(
            """CREATE TABLE IF NOT EXISTS rag_sources (
                tenant_id TEXT NOT NULL,
                source_type TEXT NOT NULL,
                source_name TEXT NOT NULL,
                content_hash TEXT NOT NULL,
                chunk_count INTEGER NOT NULL,
                updated_at TIMESTAMPTZ NOT NULL DEFAULT now(),
                PRIMARY KEY (tenant_id, source_type, source_name)
            )"""
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS rag_chunks_tsv_gin ON rag_chunks USING gin (search_tsv)"
        )
        connection.execute(
            "CREATE INDEX IF NOT EXISTS rag_chunks_embedding_hnsw "
            "ON rag_chunks USING hnsw (embedding vector_cosine_ops)"
        )
    except Exception as exc:
        log_database_error(_LOGGER, "rag_schema_initialization", exc)
        raise RagError(
            "Could not initialize the RAG schema. Confirm pgvector is enabled and the database user can create the required tables and indexes."
        ) from exc


def ensure_index() -> dict[str, Any]:
    """Check that PostgreSQL/pgvector is ready; never sends text for embedding."""
    tenant = _tenant_id()
    with _connect() as connection:
        _ensure_schema(connection)
        # Session uploads are private to one browser session and age out even
        # when their owner never returns to the app.
        connection.execute(
            "DELETE FROM rag_chunks WHERE tenant_id ~ '^public_upload_[0-9a-f]{32}$' "
            "AND updated_at < now() - interval '24 hours'"
        )
        connection.execute(
            "DELETE FROM rag_sources WHERE tenant_id ~ '^public_upload_[0-9a-f]{32}$' "
            "AND updated_at < now() - interval '24 hours'"
        )
        counts = connection.execute(
            """SELECT source_type, count(*) FROM rag_chunks
               WHERE tenant_id = %s GROUP BY source_type""",
            (tenant,),
        ).fetchall()
    by_type = {str(kind): int(count) for kind, count in counts}
    return {
        "backend": "PostgreSQL + pgvector",
        "embedding_model": _embedding_model(),
        "document_chunks": by_type.get("company_document", 0),
        "review_records": by_type.get("customer_review", 0),
        "ready": True,
    }


@trace_component("Featherless embedding request", run_type="llm")
def _embedding_request(texts: list[str]) -> list[list[float]]:
    load_local_environment()
    api_key = os.getenv("FEATHERLESS_API_KEY", "").strip()
    if not api_key:
        raise RagError("FEATHERLESS_API_KEY is missing from the project .env file.")
    base_url = os.getenv("FEATHERLESS_BASE_URL", "https://api.featherless.ai/v1").rstrip("/")
    model = _embedding_model()
    payload = {
        "model": model,
        "input": texts,
        "dimensions": EMBEDDING_DIMENSIONS,
        "encoding_format": "float",
    }
    request = Request(
        f"{base_url}/embeddings",
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "User-Agent": "BusinessAgentOS/0.1 (document embeddings)",
        },
        method="POST",
    )
    try:
        reserve_public_request()
        with urlopen(request, timeout=90) as response:
            result = json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        detail = ""
        try:
            body = json.loads(exc.read().decode("utf-8", errors="replace"))
            error = body.get("error") if isinstance(body, dict) else None
            detail = error.get("message", "") if isinstance(error, dict) else ""
        except (UnicodeDecodeError, json.JSONDecodeError, OSError):
            pass
        detail = " ".join(str(detail).split())[:240]
        suffix = f" Provider detail: {detail}" if detail else ""
        raise RagError(f"Featherless embeddings returned HTTP {exc.code}.{suffix}") from None
    except (URLError, TimeoutError, OSError):
        raise RagError("Could not reach Featherless embeddings. Try again later.") from None
    except (UnicodeDecodeError, json.JSONDecodeError):
        raise RagError("Featherless embeddings returned invalid JSON.") from None

    try:
        items = result["data"]
        vectors = [None] * len(texts)
        for item in items:
            index = int(item["index"])
            vector = [float(value) for value in item["embedding"]]
            if index < 0 or index >= len(vectors) or len(vector) != EMBEDDING_DIMENSIONS:
                raise ValueError("unexpected embedding dimensions or index")
            if not all(math.isfinite(value) for value in vector):
                raise ValueError("non-finite vector values")
            vectors[index] = vector
        if any(vector is None for vector in vectors):
            raise ValueError("missing embedding")
    except (KeyError, IndexError, TypeError, ValueError):
        raise RagError(
            f"Featherless returned embeddings that do not match the configured "
            f"{EMBEDDING_DIMENSIONS}-dimension index. Check FEATHERLESS_EMBEDDING_MODEL."
        ) from None
    return vectors


def _vector_text(vector: list[float]) -> str:
    return "[" + ",".join(format(value, ".9g") for value in vector) + "]"


def _embed_batch(texts: list[str]) -> list[list[float]]:
    vectors: list[list[float]] = []
    for start in range(0, len(texts), EMBEDDING_BATCH_SIZE):
        vectors.extend(_embedding_request(texts[start:start + EMBEDDING_BATCH_SIZE]))
    return vectors


def _chunk_text(text: str, section: str) -> list[tuple[str, str]]:
    """Split by paragraphs, retaining a small overlap and section label."""
    pieces: list[str] = []
    buffer = ""
    for paragraph in re.split(r"\n\s*\n", text):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if len(paragraph) > MAX_CHUNK_CHARS:
            if buffer:
                pieces.append(buffer)
                buffer = ""
            start = 0
            while start < len(paragraph):
                end = min(start + MAX_CHUNK_CHARS, len(paragraph))
                if end < len(paragraph):
                    boundary = paragraph.rfind(" ", start + MAX_CHUNK_CHARS // 2, end)
                    if boundary > start:
                        end = boundary
                piece = paragraph[start:end].strip()
                if piece:
                    pieces.append(piece)
                if end >= len(paragraph):
                    break
                start = max(start + 1, end - CHUNK_OVERLAP)
            continue
        candidate = f"{buffer}\n\n{paragraph}".strip() if buffer else paragraph
        if buffer and len(candidate) > MAX_CHUNK_CHARS:
            pieces.append(buffer)
            tail = buffer[-CHUNK_OVERLAP:].strip()
            buffer = f"{tail}\n\n{paragraph}".strip() if tail else paragraph
        else:
            buffer = candidate
    if buffer:
        pieces.append(buffer)
    return [(section, piece) for piece in pieces if piece]


def _text_document_chunks(source_name: str, text: str) -> list[dict[str, str]]:
    title = Path(source_name).stem.replace("_", " ").strip() or "Company document"
    chunks: list[dict[str, str]] = []
    section = title
    blocks: list[tuple[str, str]] = []
    lines: list[str] = []

    def flush() -> None:
        nonlocal lines
        body = "\n".join(lines).strip()
        if body:
            blocks.append((section, body))
        lines = []

    for line in text.splitlines():
        heading = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$", line)
        if heading:
            flush()
            section = heading.group(1).strip()
        elif not line.strip():
            flush()
        else:
            lines.append(line.rstrip())
    flush()

    for block_section, body in blocks:
        for label, content in _chunk_text(body, block_section):
            chunks.append({
                "title": title,
                "section": label,
                "content": content,
                "source_ref": "",
            })
    for number, chunk in enumerate(chunks, 1):
        chunk["source_ref"] = f"{source_name}#chunk-{number}"
    return chunks


def _pdf_chunks(source_name: str, raw: bytes) -> list[dict[str, str]]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RagError("PDF ingestion needs pypdf. Install dependencies with: py -m pip install -r requirements.txt") from exc
    try:
        reader = PdfReader(BytesIO(raw))
        title = Path(source_name).stem.replace("_", " ").strip() or "Company PDF"
        chunks: list[dict[str, str]] = []
        for page_number, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").strip()
            if not text:
                continue
            for chunk_number, (section, content) in enumerate(
                _chunk_text(text, f"Page {page_number}"), 1
            ):
                chunks.append({
                    "title": title,
                    "section": section,
                    "content": content,
                    "source_ref": f"{source_name}#page-{page_number}-chunk-{chunk_number}",
                })
        if not chunks:
            raise RagError(
                f"{Path(source_name).name} has no extractable text. Scanned PDFs need OCR before indexing."
            )
        return chunks
    except RagError:
        raise
    except Exception as exc:
        raise RagError(f"Could not extract text from PDF {Path(source_name).name}.") from exc


def _docx_chunks(source_name: str, raw: bytes) -> list[dict[str, str]]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RagError("DOCX ingestion needs python-docx. Install dependencies with: py -m pip install -r requirements.txt") from exc
    try:
        document = Document(BytesIO(raw))
        lines = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = getattr(paragraph.style, "name", "") or ""
            if style.casefold().startswith("heading"):
                level = re.search(r"(\d+)$", style)
                hashes = "#" * min(6, int(level.group(1)) if level else 2)
                lines.append(f"{hashes} {text}")
            else:
                lines.append(text)
        for table in document.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
        chunks = _text_document_chunks(source_name, "\n\n".join(lines))
        if not chunks:
            raise RagError(f"{Path(source_name).name} contains no extractable text.")
        return chunks
    except RagError:
        raise
    except Exception as exc:
        raise RagError(f"Could not extract text from DOCX {Path(source_name).name}.") from exc


def _read_file_chunks(path: Path, source_name: str) -> tuple[str, list[dict[str, str]]]:
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise RagError(f"Could not read document {path.name}.") from exc
    suffix = path.suffix.casefold()
    if suffix == ".pdf":
        chunks = _pdf_chunks(source_name, raw)
    elif suffix == ".docx":
        chunks = _docx_chunks(source_name, raw)
    else:
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise RagError(f"{path.name} is not valid UTF-8 text.") from exc
        if suffix == ".json":
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            except json.JSONDecodeError:
                pass
        chunks = _text_document_chunks(source_name, text)
    return sha256(raw).hexdigest(), chunks


def _collect_document_sources(include_uploads: bool = True) -> dict[tuple[str, str], dict[str, Any]]:
    sources: dict[tuple[str, str], dict[str, Any]] = {}
    folders = [(DEFAULT_DOCS_PATH, "company_document", "company_docs")]
    if include_uploads:
        folders.append((DEFAULT_UPLOADS_PATH, "company_document", "uploads"))
    for folder, source_type, prefix in folders:
        if not folder.is_dir():
            continue
        for path in sorted(folder.rglob("*")):
            if not path.is_file() or path.suffix.casefold() not in SUPPORTED_EXTENSIONS:
                continue
            if path.name.casefold() in EXCLUDED_DOCS:
                continue
            relative = path.relative_to(folder).as_posix()
            source_name = f"{prefix}/{relative}"
            content_hash, chunks = _read_file_chunks(path, source_name)
            if chunks:
                sources[(source_type, source_name)] = {
                    "content_hash": content_hash,
                    "chunks": chunks,
                }
    return sources


def _collect_review_sources() -> dict[tuple[str, str], dict[str, Any]]:
    sources: dict[tuple[str, str], dict[str, Any]] = {}
    if not BUSINESS_DB_PATH.is_file():
        return sources
    try:
        with sqlite3.connect(BUSINESS_DB_PATH.as_uri() + "?mode=ro", uri=True, timeout=10) as connection:
            columns = {row[1] for row in connection.execute('PRAGMA table_info("reviews")')}
            required = {"review_id", "product_id", "rating", "review_date", "review_text"}
            if not required.issubset(columns):
                return sources
            rows = connection.execute(
                "SELECT r.review_id, r.product_id, p.category, r.rating, r.review_date, r.review_text "
                "FROM reviews AS r LEFT JOIN products AS p ON p.product_id = r.product_id "
                "ORDER BY r.review_id"
            ).fetchall()
    except sqlite3.Error as exc:
        raise RagError("Could not read the demo review text for indexing.") from exc

    for review_id, product_id, category, rating, review_date, review_text in rows:
        text = str(review_text or "").strip()
        if not text:
            continue
        source_name = "business_demo.reviews"
        ref = f"reviews:{review_id}"
        category_text = str(category or "Unknown")
        content = (
            f"Product category: {category_text}. Product ID: {product_id}. "
            f"Rating: {rating}/5. Review date: {review_date}. Review text: {text}"
        )[:MAX_CHUNK_CHARS]
        sources[("customer_review", f"{source_name}/{ref}")] = {
            "content_hash": sha256(content.encode("utf-8")).hexdigest(),
            "chunks": [{
                "title": f"Customer review {review_id}",
                "section": f"{category_text} category; product {product_id}; rating {rating}/5; date {review_date}",
                "content": content,
                "source_ref": f"{source_name}#{ref}",
            }],
            "source_name": source_name,
            "source_ref": f"{source_name}#{ref}",
        }
    return sources


def _embedding_text(chunk: dict[str, str], source_type: str) -> str:
    prefix = (
        "Represent this customer review for a business analyst. "
        if source_type == "customer_review"
        else "Represent this company document passage for business question retrieval. "
    )
    return f"{prefix}Title: {chunk['title']}\nSection: {chunk['section']}\nPassage: {chunk['content']}"


def _uploaded_document_sources(uploaded_files: Iterable[Any]) -> tuple[dict[tuple[str, str], dict[str, Any]], int]:
    """Parse public session uploads in memory; never write them to shared disk."""
    sources: dict[tuple[str, str], dict[str, Any]] = {}
    total_bytes = 0
    files = list(uploaded_files)
    if len(files) > 3:
        raise RagError("Attach no more than three documents at a time in the public demo.")
    for uploaded in files:
        name = Path(str(getattr(uploaded, "name", "upload.pdf"))).name
        suffix = Path(name).suffix.casefold()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise RagError(f"{name} is not a supported document type.")
        raw = uploaded.getvalue()
        if not isinstance(raw, bytes) or not raw:
            continue
        total_bytes += len(raw)
        if len(raw) > 5 * 1024 * 1024 or total_bytes > 10 * 1024 * 1024:
            raise RagError("Public demo uploads are limited to 5 MB per file and 10 MB per request.")
        safe_name = re.sub(r"[^A-Za-z0-9._ -]+", "_", name).strip(" .")
        safe_name = safe_name or ("uploaded_document" + suffix)
        if suffix == ".pdf":
            chunks = _pdf_chunks(safe_name, raw)
        elif suffix == ".docx":
            chunks = _docx_chunks(safe_name, raw)
        else:
            try:
                text = raw.decode("utf-8")
            except UnicodeDecodeError as exc:
                raise RagError(f"{safe_name} is not valid UTF-8 text.") from exc
            if suffix == ".json":
                try:
                    text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
                except json.JSONDecodeError:
                    pass
            chunks = _text_document_chunks(safe_name, text)
        if not chunks:
            raise RagError(f"{safe_name} contains no extractable text.")
        source_name = f"session_uploads/{safe_name}"
        sources[("company_document", source_name)] = {
            "content_hash": sha256(raw).hexdigest(),
            "chunks": chunks,
        }
    return sources, len(sources)


def _persist_uploaded_files(uploaded_files: Iterable[Any]) -> int:
    saved = 0
    for uploaded in uploaded_files:
        name = Path(str(getattr(uploaded, "name", "upload.txt"))).name
        suffix = Path(name).suffix.casefold()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise RagError(f"{name} is not a supported document type.")
        data = uploaded.getvalue()
        if not isinstance(data, bytes) or not data:
            continue
        if len(data) > 25 * 1024 * 1024:
            raise RagError(f"{name} is larger than the 25 MB upload limit.")
        safe_name = re.sub(r"[^A-Za-z0-9._ -]+", "_", name).strip(" .")
        if not safe_name:
            safe_name = "uploaded_document" + suffix
        DEFAULT_UPLOADS_PATH.mkdir(parents=True, exist_ok=True)
        target = DEFAULT_UPLOADS_PATH / safe_name
        try:
            target.write_bytes(data)
        except OSError as exc:
            raise RagError(f"Could not save uploaded file {safe_name}.") from exc
        saved += 1
    return saved


def ingest_company_sources(
    uploaded_files: Iterable[Any] = (), *, include_uploads: bool = True,
    include_demo_sources: bool = True, tenant_id: str | None = None,
    prune_missing: bool = True,
) -> dict[str, Any]:
    """Index bundled sources and optionally local or session-isolated uploads."""
    files = list(uploaded_files)
    if tenant_id is not None:
        sources, uploaded_count = _uploaded_document_sources(files)
    else:
        uploaded_count = _persist_uploaded_files(files) if include_uploads else 0
        sources = {}
    if include_demo_sources:
        sources.update(_collect_document_sources(include_uploads=include_uploads and tenant_id is None))
        sources.update(_collect_review_sources())
    tenant = _tenant_id(tenant_id)

    result = {
        "new_sources": 0,
        "updated_sources": 0,
        "unchanged_sources": 0,
        "removed_sources": 0,
        "document_chunks": 0,
        "review_records": 0,
        "uploaded_files": uploaded_count,
        "embedding_model": _embedding_model(),
    }
    with _connect() as connection:
        _ensure_schema(connection)
        manifest_rows = connection.execute(
            """SELECT source_type, source_name, content_hash FROM rag_sources
               WHERE tenant_id = %s""",
            (tenant,),
        ).fetchall()
        review_rows = connection.execute(
            """SELECT source_name, source_ref, content_hash FROM rag_chunks
               WHERE tenant_id = %s AND source_type = 'customer_review'""",
            (tenant,),
        ).fetchall()
    manifests = {(row[0], row[1]): row[2] for row in manifest_rows}
    review_hashes = {(row[0], row[1]): row[2] for row in review_rows}

    pending: list[dict[str, Any]] = []
    current_document_names: set[str] = set()
    current_review_refs: set[str] = set()
    for (source_type, source_key), source in sources.items():
        if source_type == "customer_review":
            source_name = str(source["source_name"])
            source_ref = str(source["source_ref"])
            current_review_refs.add(source_ref)
            prior_hash = review_hashes.get((source_name, source_ref))
            if prior_hash == source["content_hash"]:
                result["unchanged_sources"] += 1
                continue
            pending.append({
                "source_type": source_type,
                "source_name": source_name,
                "source_ref": source_ref,
                "content_hash": source["content_hash"],
                "chunks": source["chunks"],
                "was_present": prior_hash is not None,
            })
            continue

        source_name = source_key
        current_document_names.add(source_name)
        prior_hash = manifests.get((source_type, source_name))
        chunks = source["chunks"]
        if prior_hash == source["content_hash"]:
            result["unchanged_sources"] += 1
            result["document_chunks"] += len(chunks)
            continue
        pending.append({
            "source_type": source_type,
            "source_name": source_name,
            "content_hash": source["content_hash"],
            "chunks": chunks,
            "was_present": prior_hash is not None,
        })

    flat_texts = [
        _embedding_text(chunk, item["source_type"])
        for item in pending
        for chunk in item["chunks"]
    ]
    vectors = _embed_batch(flat_texts) if flat_texts else []
    offset = 0
    prepared: list[tuple[dict[str, Any], list[list[float]]]] = []
    for item in pending:
        count = len(item["chunks"])
        prepared.append((item, vectors[offset:offset + count]))
        offset += count

    with _connect() as connection:
        _ensure_schema(connection)
        for item, item_vectors in prepared:
            source_type = item["source_type"]
            source_name = item["source_name"]
            chunks = item["chunks"]
            if source_type == "customer_review":
                chunk = chunks[0]
                connection.execute(
                    """INSERT INTO rag_chunks
                       (tenant_id, source_type, source_name, source_ref, title, section,
                        content, content_hash, embedding)
                       VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s::vector)
                       ON CONFLICT (tenant_id, source_name, source_ref) DO UPDATE SET
                         source_type = EXCLUDED.source_type, title = EXCLUDED.title,
                         section = EXCLUDED.section, content = EXCLUDED.content,
                         content_hash = EXCLUDED.content_hash, embedding = EXCLUDED.embedding,
                         updated_at = now()""",
                    (
                        tenant, source_type, source_name, item["source_ref"],
                        chunk["title"], chunk["section"], chunk["content"],
                        item["content_hash"], _vector_text(item_vectors[0]),
                    ),
                )
                result["review_records"] += 1
            else:
                connection.execute(
                    "DELETE FROM rag_chunks WHERE tenant_id = %s AND source_name = %s",
                    (tenant, source_name),
                )
                for chunk, vector in zip(chunks, item_vectors):
                    content_hash_chunk = sha256(chunk["content"].encode("utf-8")).hexdigest()
                    connection.execute(
                        """INSERT INTO rag_chunks
                           (tenant_id, source_type, source_name, source_ref, title, section,
                            content, content_hash, embedding)
                           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s::vector)""",
                        (
                            tenant, source_type, source_name, chunk["source_ref"],
                            chunk["title"], chunk["section"], chunk["content"],
                            content_hash_chunk, _vector_text(vector),
                        ),
                    )
                connection.execute(
                    """INSERT INTO rag_sources
                       (tenant_id, source_type, source_name, content_hash, chunk_count, updated_at)
                       VALUES (%s, %s, %s, %s, %s, now())
                       ON CONFLICT (tenant_id, source_type, source_name) DO UPDATE SET
                         content_hash = EXCLUDED.content_hash, chunk_count = EXCLUDED.chunk_count,
                         updated_at = now()""",
                    (tenant, source_type, source_name, item["content_hash"], len(chunks)),
                )
                result["document_chunks"] += len(chunks)

            result["updated_sources" if item["was_present"] else "new_sources"] += 1

        for old_type, old_name, _old_hash in manifest_rows:
            if prune_missing and old_type == "company_document" and old_name not in current_document_names:
                connection.execute(
                    "DELETE FROM rag_chunks WHERE tenant_id = %s AND source_name = %s",
                    (tenant, old_name),
                )
                connection.execute(
                    "DELETE FROM rag_sources WHERE tenant_id = %s AND source_type = %s AND source_name = %s",
                    (tenant, old_type, old_name),
                )
                result["removed_sources"] += 1

        stored_reviews = connection.execute(
            """SELECT source_ref FROM rag_chunks WHERE tenant_id = %s
               AND source_type = 'customer_review'""",
            (tenant,),
        ).fetchall()
        for (old_ref,) in stored_reviews:
            if prune_missing and old_ref not in current_review_refs:
                connection.execute(
                    "DELETE FROM rag_chunks WHERE tenant_id = %s AND source_ref = %s",
                    (tenant, old_ref),
                )
                result["removed_sources"] += 1

    return result

def _query_embedding(query: str) -> list[float]:
    instruction = (
        "Instruct: Given a business question, retrieve passages that are useful "
        "to answer it.\nQuery: "
    )
    return _embedding_request([instruction + query])[0]


def _requested_product_categories(query: str) -> list[str]:
    """Find exact product-category mentions so review search can respect them."""
    if not BUSINESS_DB_PATH.is_file():
        return []
    try:
        with sqlite3.connect(BUSINESS_DB_PATH.as_uri() + "?mode=ro", uri=True, timeout=5) as connection:
            categories = [str(row[0]) for row in connection.execute(
                "SELECT DISTINCT category FROM products WHERE category IS NOT NULL"
            )]
    except sqlite3.Error:
        return []
    return [category for category in categories
            if re.search(r"(?<!\w)" + re.escape(category) + r"(?!\w)", query, re.IGNORECASE)]


@trace_component("Semantic document retrieval", run_type="retriever")
def retrieve_relevant(
    query: str,
    top_k: int = 5,
    tenant_id: str | list[str] | tuple[str, ...] | None = None,
) -> list[dict[str, Any]]:
    """Hybrid retrieval over the requested tenant(s), without crossing tenant boundaries."""
    if not isinstance(query, str) or not query.strip():
        return []
    if tenant_id is None:
        active_tenants = _ACTIVE_RETRIEVAL_TENANTS.get()
        tenants = list(active_tenants) if active_tenants else [_tenant_id()]
    elif isinstance(tenant_id, str):
        tenants = [tenant_id.strip()]
    else:
        tenants = [value.strip() for value in tenant_id if isinstance(value, str)]
    tenants = list(dict.fromkeys(value for value in tenants if value))
    if not tenants or any(len(value) > 128 for value in tenants):
        raise RagError("The retrieval tenant identifier is invalid.")
    limit = max(1, min(int(top_k), 10))
    with _connect() as connection:
        _ensure_schema(connection)
        total = connection.execute(
            "SELECT count(*) FROM rag_chunks WHERE tenant_id = ANY(%s)",
            (tenants,),
        ).fetchone()[0]
        if not total:
            raise RagError("The semantic index is empty. Index company documents in the sidebar first.")

        # Check the local index before transmitting a query for embedding.
        query_vector = _vector_text(_query_embedding(query.strip()))
        semantic = connection.execute(
            """SELECT id, source_type, source_name, source_ref, title, section, content,
                      1 - (embedding <=> %s::vector) AS score
               FROM rag_chunks
               WHERE tenant_id = ANY(%s)
               ORDER BY embedding <=> %s::vector
               LIMIT 25""",
            (query_vector, tenants, query_vector),
        ).fetchall()
        lexical_terms = [token for token in re.findall(r"[A-Za-z0-9]+", query) if len(token) > 1]
        lexical_query = " | ".join(lexical_terms)
        lexical = connection.execute(
            """SELECT id, source_type, source_name, source_ref, title, section, content,
                      ts_rank_cd(search_tsv, to_tsquery('english', %s)) AS score
               FROM rag_chunks
               WHERE tenant_id = ANY(%s)
                 AND search_tsv @@ to_tsquery('english', %s)
               ORDER BY score DESC
               LIMIT 25""",
            (lexical_query, tenants, lexical_query),
        ).fetchall() if lexical_query else []

        # When both review polarities are requested, retrieve examples from each
        # rating band instead of allowing many near-duplicate positive reviews to
        # crowd out all critical reviews. The actual review text remains evidence.
        query_lower = query.casefold()
        wants_both_polarities = (
            re.search(r"\b(positive|positives|favorable|favourable)\b", query_lower)
            and re.search(r"\b(negative|negatives|critical|complaints?)\b", query_lower)
            and re.search(r"\b(reviews?|sentiment|themes?)\b", query_lower)
        )
        polarity_rows: dict[str, list[Any]] = {"positive": [], "negative": []}
        if wants_both_polarities:
            categories = _requested_product_categories(query)
            category_clause = ""
            category_params: list[str] = []
            if categories:
                category_clause = " AND (" + " OR ".join("section ILIKE %s" for _ in categories) + ")"
                category_params.extend(f"%{category} category%" for category in categories)
            for polarity, ratings in (("negative", (1, 2)), ("positive", (4, 5))):
                rating_clause = "(" + " OR ".join("section ILIKE %s" for _ in ratings) + ")"
                rating_params = [f"%rating {rating}/5%" for rating in ratings]
                rows = connection.execute(
                    """SELECT id, source_type, source_name, source_ref, title, section, content,
                              1 - (embedding <=> %s::vector) AS score
                       FROM rag_chunks
                       WHERE tenant_id = ANY(%s) AND source_type = 'customer_review'"""
                    + category_clause + " AND " + rating_clause
                    + " ORDER BY embedding <=> %s::vector, source_ref LIMIT 100",
                    (query_vector, tenants, *category_params, *rating_params, query_vector),
                ).fetchall()
                # Repeated synthetic/customer phrases otherwise occupy most of
                # the nearest-neighbor slots. Keep the nearest distinct text themes.
                unique_rows = []
                seen_themes: set[str] = set()
                for row in rows:
                    theme_match = re.search(r"Review text:\s*(.*)$", str(row[6]))
                    theme = theme_match.group(1).strip() if theme_match else str(row[6])
                    if theme in seen_themes:
                        continue
                    seen_themes.add(theme)
                    unique_rows.append(row)
                    if len(unique_rows) >= limit:
                        break
                polarity_rows[polarity] = unique_rows

    combined: dict[int, dict[str, Any]] = {}
    rrf_scores: defaultdict[int, float] = defaultdict(float)
    for rank, row in enumerate(semantic, 1):
        item_id = int(row[0])
        rrf_scores[item_id] += 1 / (60 + rank)
        combined[item_id] = {
            "source_type": row[1], "source_name": row[2], "source_ref": row[3],
            "title": row[4], "section": row[5], "content": row[6],
            "semantic_score": float(row[7]), "retrieval_method": "hybrid_semantic",
        }
    for rank, row in enumerate(lexical, 1):
        item_id = int(row[0])
        rrf_scores[item_id] += 1 / (60 + rank)
        if item_id not in combined:
            combined[item_id] = {
                "source_type": row[1], "source_name": row[2], "source_ref": row[3],
                "title": row[4], "section": row[5], "content": row[6],
                "semantic_score": None, "retrieval_method": "hybrid_keyword",
            }
        else:
            combined[item_id]["retrieval_method"] = "hybrid_semantic+keyword"
    ranked = sorted(combined, key=lambda item_id: (-rrf_scores[item_id], item_id))
    if wants_both_polarities:
        diverse: list[int] = []
        # Alternate distinct themes from both rating bands, then fill from the
        # ordinary hybrid ranking without repeating the same review wording.
        for theme_rank in range(limit):
            for polarity in ("negative", "positive"):
                if theme_rank >= len(polarity_rows[polarity]):
                    continue
                row = polarity_rows[polarity][theme_rank]
                item_id = int(row[0])
                rrf_scores[item_id] += 1 / (60 + theme_rank + 1)
                combined[item_id] = {
                    "source_type": row[1], "source_name": row[2], "source_ref": row[3],
                    "title": row[4], "section": row[5], "content": row[6],
                    "semantic_score": float(row[7]), "retrieval_method": "hybrid_semantic+rating_diversity",
                }
                if item_id not in diverse:
                    diverse.append(item_id)
        seen_review_themes = {
            match.group(1).strip()
            for item_id in diverse
            if (match := re.search(r"Review text:\s*(.*)$", str(combined[item_id].get("content", ""))))
        }
        remainder = []
        for item_id in ranked:
            if item_id in diverse:
                continue
            item = combined[item_id]
            if item.get("source_type") == "customer_review":
                match = re.search(r"Review text:\s*(.*)$", str(item.get("content", "")))
                theme = match.group(1).strip() if match else str(item.get("content", ""))
                if theme in seen_review_themes:
                    continue
                seen_review_themes.add(theme)
            remainder.append(item_id)
        ranked = diverse + remainder
    return [
        {
            **combined[item_id],
            "score": float(rrf_scores[item_id]),
        }
        for item_id in ranked[:limit]
    ]

def ingest_uploaded_files(
    uploaded_files: Iterable[Any] = (), **options: Any
) -> dict[str, Any]:
    """UI alias; local uploads persist, public uploads stay in their session tenant."""
    return ingest_company_sources(uploaded_files=uploaded_files, **options)


